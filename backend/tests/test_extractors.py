"""Extractor tests: each supported format plus unknown-file fallback."""
import io

import pytest

from app.ingest.extractors import (ExtractionError, UnsupportedFormat,
                                   extract)


def test_txt():
    units = extract(b"hello world", "note.txt")
    assert units[0].text == "hello world"


def test_markdown_sections():
    md = b"# Title\nintro\n\n## Details\nbody text"
    units = extract(md, "doc.md")
    assert any(u.section == "Details" for u in units)
    assert all(u.text for u in units)


def test_csv_row_provenance():
    rows = b"name,age\nalice,30\nbob,40\n"
    units = extract(rows, "people.csv")
    assert units[0].meta["row_start"] == 1
    assert "alice" in units[0].text


def test_html_strips_scripts():
    html = b"<html><head><script>x=1</script></head><body><p>Hi there</p></body></html>"
    units = extract(html, "page.html")
    assert "Hi there" in units[0].text
    assert "x=1" not in units[0].text


def test_pdf_pages():
    from pypdf import PdfWriter

    buf = io.BytesIO()
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    w.write(buf)
    # A blank page yields no text -> ExtractionError is the honest result.
    with pytest.raises(ExtractionError):
        extract(buf.getvalue(), "blank.pdf")


def test_docx_paragraphs_and_tables():
    from docx import Document

    doc = Document()
    doc.add_heading("Contract", level=1)
    doc.add_paragraph("Payment due in 30 days.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Price"
    buf = io.BytesIO()
    doc.save(buf)
    units = extract(buf.getvalue(), "contract.docx")
    texts = "\n".join(u.text for u in units)
    assert "Payment due in 30 days." in texts
    assert "Item | Price" in texts
    assert any(u.section == "Contract" for u in units)


def test_xlsx_sheet_provenance():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Cost"])
    ws.append(["Laptop", 1200])
    buf = io.BytesIO()
    wb.save(buf)
    units = extract(buf.getvalue(), "budget.xlsx")
    assert units[0].sheet == "Budget"
    assert "Laptop" in units[0].text


def test_pptx_slide_provenance():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Quarterly results"
    buf = io.BytesIO()
    prs.save(buf)
    units = extract(buf.getvalue(), "deck.pptx")
    assert units[0].slide == 1
    assert "Quarterly results" in units[0].text


def test_unknown_binary_rejected():
    with pytest.raises(UnsupportedFormat):
        extract(b"\x00\x01\x02\xff\xfe binary", "mystery.bin")


def test_unknown_text_extension_falls_back():
    units = extract("plain unicode text ✓".encode("utf-8"), "notes.weird")
    assert "unicode text" in units[0].text


def test_extensionless_office_detected_by_content():
    from docx import Document

    doc = Document()
    doc.add_paragraph("hidden docx")
    buf = io.BytesIO()
    doc.save(buf)
    units = extract(buf.getvalue(), "noext")
    assert "hidden docx" in units[0].text
