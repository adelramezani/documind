"""Document text extraction registry.

Each extractor turns raw file bytes into a list of ``ExtractedUnit`` —
a piece of text plus provenance metadata (page, sheet, slide, heading).
Unknown extensions fall back to safe text-sniffing; binary/unsupported
files raise ``UnsupportedFormat`` and are reported to the user, never
executed or handed to external services.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Callable

from charset_normalizer import from_bytes


class UnsupportedFormat(Exception):
    """Raised when no extractor can handle the file."""


class ExtractionError(Exception):
    """Raised when a registered extractor fails on the content."""


@dataclass
class ExtractedUnit:
    text: str
    page: int | None = None
    sheet: str | None = None
    slide: int | None = None
    section: str | None = None
    meta: dict = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Plain text-like formats
# ---------------------------------------------------------------------------

def _decode_text(data: bytes) -> str:
    match = from_bytes(data).best()
    if match is None:
        raise ExtractionError("could not decode text content")
    return str(match)


def extract_text(data: bytes, name: str) -> list[ExtractedUnit]:
    return [ExtractedUnit(text=_decode_text(data))]


def extract_markdown(data: bytes, name: str) -> list[ExtractedUnit]:
    text = _decode_text(data)
    # Split by headings so provenance can record the section.
    units: list[ExtractedUnit] = []
    current_section: str | None = None
    buf: list[str] = []
    for line in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            if buf:
                units.append(ExtractedUnit("\n".join(buf), section=current_section))
                buf = []
            current_section = m.group(2).strip()
        buf.append(line)
    if buf:
        units.append(ExtractedUnit("\n".join(buf), section=current_section))
    return units or [ExtractedUnit(text=text)]


def extract_csv(data: bytes, name: str) -> list[ExtractedUnit]:
    text = _decode_text(data)
    rows = list(csv.reader(io.StringIO(text)))
    units: list[ExtractedUnit] = []
    header = rows[0] if rows else []
    batch, start = [], 0
    for i, row in enumerate(rows):
        line = " | ".join(row)
        batch.append(line)
        if len(batch) >= 50:
            units.append(
                ExtractedUnit(
                    "\n".join(batch),
                    meta={"row_start": start + 1, "row_end": i + 1,
                          "header": header},
                )
            )
            start = i + 1
            batch = []
    if batch:
        units.append(
            ExtractedUnit("\n".join(batch),
                          meta={"row_start": start + 1,
                                "row_end": len(rows), "header": header})
        )
    return units or [ExtractedUnit(text=text)]


def extract_html(data: bytes, name: str) -> list[ExtractedUnit]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode_text(data), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = soup.get_text("\n")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return [ExtractedUnit(text=text)]


# ---------------------------------------------------------------------------
# Binary Office / PDF formats
# ---------------------------------------------------------------------------

def extract_pdf(data: bytes, name: str) -> list[ExtractedUnit]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - pypdf raises various errors
        raise ExtractionError(f"could not read PDF: {exc}") from exc
    units = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if text:
            units.append(ExtractedUnit(text=text, page=i + 1))
    if not units:
        raise ExtractionError("PDF contains no extractable text (scanned?)")
    return units


def extract_docx(data: bytes, name: str) -> list[ExtractedUnit]:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"could not read DOCX: {exc}") from exc
    units: list[ExtractedUnit] = []
    buf: list[str] = []
    section: str | None = None
    for para in doc.paragraphs:
        style = (para.style.name or "") if para.style else ""
        if style.startswith("Heading"):
            if buf:
                units.append(ExtractedUnit("\n".join(buf), section=section))
                buf = []
            section = para.text.strip()
        elif para.text.strip():
            buf.append(para.text)
    if buf:
        units.append(ExtractedUnit("\n".join(buf), section=section))
    for table in doc.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows]
        if rows:
            units.append(ExtractedUnit("\n".join(rows),
                                       section=section,
                                       meta={"table": True}))
    if not units:
        raise ExtractionError("DOCX contains no extractable text")
    return units


def extract_xlsx(data: bytes, name: str) -> list[ExtractedUnit]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"could not read XLSX: {exc}") from exc
    units: list[ExtractedUnit] = []
    for ws in wb.worksheets:
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            for start in range(0, len(rows), 50):
                units.append(
                    ExtractedUnit(
                        "\n".join(rows[start:start + 50]),
                        sheet=ws.title,
                        meta={"row_start": start + 1,
                              "row_end": min(start + 50, len(rows))},
                    )
                )
    if not units:
        raise ExtractionError("workbook contains no data")
    return units


def extract_pptx(data: bytes, name: str) -> list[ExtractedUnit]:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"could not read PPTX: {exc}") from exc
    units = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text for c in row.cells))
        if texts:
            units.append(ExtractedUnit("\n".join(texts), slide=i + 1))
    if not units:
        raise ExtractionError("PPTX contains no extractable text")
    return units


# ---------------------------------------------------------------------------
# Registry & fallback
# ---------------------------------------------------------------------------

EXTENSION_MAP: dict[str, Callable[[bytes, str], list[ExtractedUnit]]] = {
    ".txt": extract_text,
    ".log": extract_text,
    ".md": extract_markdown,
    ".markdown": extract_markdown,
    ".csv": extract_csv,
    ".tsv": extract_csv,
    ".html": extract_html,
    ".htm": extract_html,
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".xlsm": extract_xlsx,
    ".pptx": extract_pptx,
}

_TEXT_SNIFF_BYTES = 4096


def _looks_like_text(data: bytes) -> bool:
    sample = data[:_TEXT_SNIFF_BYTES]
    if b"\x00" in sample:
        return False
    try:
        sample.decode("utf-8")
        return True
    except UnicodeDecodeError:
        return False


def extract(data: bytes, filename: str) -> list[ExtractedUnit]:
    """Extract text units from file bytes, with safe fallback for unknown
    extensions (text sniffing only — never executes content)."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    fn = EXTENSION_MAP.get(ext)
    if fn is not None:
        return fn(data, filename)
    # Unknown extension: try safe text extraction.
    if _looks_like_text(data):
        return extract_text(data, filename)
    # Office formats are zip-based: detect docx/xlsx/pptx by content.
    if zipfile.is_zipfile(io.BytesIO(data)):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
            if "word/document.xml" in names:
                return extract_docx(data, filename)
            if "xl/workbook.xml" in names:
                return extract_xlsx(data, filename)
            if "ppt/presentation.xml" in names:
                return extract_pptx(data, filename)
        except (ExtractionError, zipfile.BadZipFile):
            pass
    raise UnsupportedFormat(
        f"unsupported file type '{ext or 'unknown'}' for {filename!r}"
    )


SUPPORTED_EXTENSIONS = sorted(EXTENSION_MAP)
